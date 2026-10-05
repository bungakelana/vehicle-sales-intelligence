"""Generate dummy raw sales data (1000 rows) following the 36 columns in PROJECT_BRIEF section 3.

Data meniru dealer kendaraan niaga di Riau & sekitarnya, lengkap dengan "kotoran" khas data mentah
(penulisan kota tidak konsisten, leasing kosong vs "CASH", tanggal proses belum terisi untuk unit baru)
supaya Fase 1-2 (audit & cleaning) punya sesuatu untuk dikerjakan.

Jalankan: python generate_dummy_data.py
"""
from datetime import date, timedelta
from pathlib import Path
import random

import numpy as np
import pandas as pd
from faker import Faker
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEED = 42
N_ROWS = 1000
START, END = date(2024, 1, 2), date(2026, 9, 30)
OUT = Path(__file__).parent / "analytics" / "data" / "raw" / "dummy_vehicle_sales.xlsx"

rng = np.random.default_rng(SEED)
random.seed(SEED)
fake = Faker("id_ID")
Faker.seed(SEED)

# --- Geografi: kabupaten -> (bobot, varian penulisan, sales point terdekat, kode plat, jarak hari kirim)
CITIES = {
    "Pekanbaru":        (0.24, ["PEKANBARU", "Kota Pekanbaru", "PKU", "Pekanbaru"], "PEKANBARU", "BM", (1, 4)),
    "Kampar":           (0.13, ["KAMPAR", "Kab. Kampar", "BANGKINANG"], "BANGKINANG", "BM", (2, 6)),
    "Siak":             (0.08, ["SIAK", "Kab Siak", "SIAK SRI INDRAPURA"], "PEKANBARU", "BM", (3, 7)),
    "Pelalawan":        (0.09, ["PELALAWAN", "PANGKALAN KERINCI", "Kab. Pelalawan"], "PEKANBARU", "BM", (3, 8)),
    "Bengkalis":        (0.08, ["BENGKALIS", "DURI", "Kab. Bengkalis", "MANDAU"], "DURI", "BM", (2, 7)),
    "Dumai":            (0.07, ["DUMAI", "Kota Dumai"], "DUMAI", "BM", (1, 4)),
    "Rokan Hulu":       (0.07, ["ROKAN HULU", "PASIR PENGARAIAN", "ROHUL"], "BANGKINANG", "BM", (4, 10)),
    "Rokan Hilir":      (0.05, ["ROKAN HILIR", "BAGAN SIAPIAPI", "ROHIL"], "DUMAI", "BM", (4, 10)),
    "Indragiri Hulu":   (0.06, ["INDRAGIRI HULU", "RENGAT", "INHU"], "RENGAT", "BM", (2, 6)),
    "Indragiri Hilir":  (0.04, ["INDRAGIRI HILIR", "TEMBILAHAN", "INHIL"], "RENGAT", "BM", (5, 14)),
    "Kuantan Singingi": (0.04, ["KUANTAN SINGINGI", "TELUK KUANTAN", "KUANSING"], "RENGAT", "BM", (3, 8)),
    "Kepulauan Meranti":(0.01, ["KEPULAUAN MERANTI", "SELATPANJANG"], "DUMAI", "BM", (7, 16)),
    "Padang":           (0.02, ["PADANG", "Kota Padang"], "PEKANBARU", "BA", (5, 12)),
    "Jambi":            (0.01, ["JAMBI", "Kota Jambi"], "RENGAT", "BH", (5, 12)),
    "Medan":            (0.01, ["MEDAN"], "DUMAI", "BK", (6, 14)),
}
MINING_PLANTATION = {"Kampar", "Pelalawan", "Rokan Hulu", "Rokan Hilir", "Indragiri Hulu", "Kuantan Singingi", "Siak"}
SALES_POINTS = ["PEKANBARU", "BANGKINANG", "DURI", "DUMAI", "RENGAT"]

# --- Produk: family -> (bobot, list (spec_no, harga dasar 2024 dalam juta), karoseri yang mungkin, base diskon %)
PRODUCTS = {
    "PICKUP":        (0.30, [("PU-150-STD", 238), ("PU-150-DLX", 262), ("PU-150-DC", 289)],
                      ["BAK BESI", "BAK KAYU", "BOX ALUMINIUM", "CAB CHASSIS"], 3.0),
    "LIGHT TRUCK":   (0.30, [("LT-110-PS", 336), ("LT-110-HD", 358), ("LT-125-HD", 381)],
                      ["BAK KAYU", "BOX BESI", "BOX ALUMINIUM", "DUMP", "CAB CHASSIS"], 4.0),
    "LIGHT TRUCK 6": (0.20, [("LTD-136-HD", 448), ("LTD-150-LONG", 497)],
                      ["DUMP", "BAK KAYU", "BOX BESI", "TANGKI", "CAB CHASSIS"], 4.5),
    "MEDIUM TRUCK":  (0.14, [("MT-190-4X2", 712), ("MT-235-6X2", 838)],
                      ["DUMP", "TANGKI CPO", "WING BOX", "BAK KAYU", "CAB CHASSIS"], 5.0),
    "HEAVY TRUCK":   (0.06, [("HT-280-6X4", 1285), ("HT-340-TRACTOR", 1520)],
                      ["DUMP", "TANGKI CPO", "MIXER", "CAB CHASSIS"], 5.5),
}
COLORS = ["PUTIH", "PUTIH", "PUTIH", "PUTIH", "SILVER", "MERAH", "BIRU", "KUNING", "HIJAU", "HITAM"]

# Leasing: penulisan sengaja tidak konsisten (untuk latihan standarisasi)
LEASING = [
    (["NUSANTARA FINANCE", "PT NUSANTARA FINANCE", "NUSANTARA FIN"], 0.32),
    (["SINAR MITRA MULTIFINANCE", "SINAR MITRA MF"], 0.24),
    (["BINA ARTHA LEASING", "PT. BINA ARTHA LEASING"], 0.19),
    (["MEGA DANA FINANCE"], 0.10),
    (["ANDALAS MULTIFINANCE"], 0.08),
    (["CAKRAWALA FINANCE"], 0.07),
]
PACKAGES = ["PKG-REGULER", "PKG-REGULER", "PKG-FLEET", "PKG-DP-RINGAN", "PKG-SERVICE-FREE", "PKG-AKHIR-TAHUN"]

# --- Tim sales: 5 SPV, 22 salesman; tiap salesman punya bias diskon & produktivitas
SPVS = {sp: fake.name().upper() for sp in SALES_POINTS}
SALESMEN = []
for i in range(22):
    sp = SALES_POINTS[i % len(SALES_POINTS)]
    SALESMEN.append({
        "name": fake.name().upper(),
        "sales_point": sp,
        "spv": SPVS[sp],
        "weight": rng.gamma(2.0, 1.0),
        "disc_bias": rng.normal(0, 1.0) + (2.5 if i in (3, 11) else 0),  # 2 salesman "discount-driven"
    })


def pick(weights_dict):
    keys = list(weights_dict)
    w = np.array([weights_dict[k][0] for k in keys], dtype=float)
    return keys[rng.choice(len(keys), p=w / w.sum())]


def random_billing_date():
    """Tren naik, musiman (dip saat Lebaran, puncak Des), dan lonjakan akhir bulan."""
    while True:
        d = START + timedelta(days=int(rng.integers(0, (END - START).days + 1)))
        if d.weekday() == 6:
            continue
        months_from_start = (d.year - START.year) * 12 + d.month - 1
        w = 1 + 0.012 * months_from_start
        if (d.year, d.month) in {(2024, 4), (2025, 3), (2026, 3)}:  # sekitar Lebaran
            w *= 0.7
        if d.month in (11, 12):
            w *= 1.3
        if d.day >= 24:
            w *= 1.6
        if rng.random() < w / 3.0:
            return d


# --- Customer pool: retail + fleet/corporate (fleet beli berulang)
def make_customer(i):
    corporate = rng.random() < 0.35
    if corporate:
        prefix = random.choice(["PT", "PT", "PT.", "CV", "CV.", "UD", "KOPERASI"])
        name = f"{prefix} {fake.last_name().upper()} {random.choice(['JAYA', 'ABADI', 'SAWIT LESTARI', 'MANDIRI', 'TRANS LOGISTIK', 'AGRO', 'SEJAHTERA', 'KARYA'])}"
    else:
        name = fake.name().upper()
    kab = pick(CITIES)
    return {
        "id": f"C{1000000 + i:07d}",
        "name": name,
        "corporate": corporate,
        "kab": kab,
        "city_raw": random.choice(CITIES[kab][1]),
        "address": fake.street_address().upper(),
        "hp": fake.numerify("08##########"),
        "home": fake.numerify("0761-######") if rng.random() < 0.25 else None,
        "office": fake.numerify("0761-######") if corporate or rng.random() < 0.1 else None,
        "weight": rng.pareto(1.2) + 1 if corporate else 1.0,
    }


customers = [make_customer(i) for i in range(620)]
cust_w = np.array([c["weight"] for c in customers])
cust_w /= cust_w.sum()

rows = []
used_plates = set()
for _ in range(N_ROWS):
    bdate = random_billing_date()
    cust = customers[rng.choice(len(customers), p=cust_w)]
    kab = cust["kab"]
    _, _, near_sp, plate_code, ship_days = CITIES[kab]

    # salesman: 75% dari sales point terdekat
    pool = [s for s in SALESMEN if s["sales_point"] == near_sp] if rng.random() < 0.75 else SALESMEN
    sw = np.array([s["weight"] for s in pool])
    sm = pool[rng.choice(len(pool), p=sw / sw.sum())]

    # produk: wilayah tambang/perkebunan lebih banyak truk besar
    fam_w = {k: (v[0] * (1.8 if kab in MINING_PLANTATION and k in ("MEDIUM TRUCK", "HEAVY TRUCK", "LIGHT TRUCK 6") else 1),) for k, v in PRODUCTS.items()}
    fam = pick(fam_w)
    _, specs, bodies, base_disc = PRODUCTS[fam]
    spec, base_price = random.choice(specs)
    if kab in MINING_PLANTATION and "DUMP" in bodies and rng.random() < 0.4:
        body = "DUMP"
    elif kab in ("Pekanbaru", "Dumai") and rng.random() < 0.3 and "BOX ALUMINIUM" in bodies:
        body = "BOX ALUMINIUM"
    else:
        body = random.choice(bodies)

    # harga: naik ~3,5%/tahun, dibulatkan ke 500rb
    years = (bdate - START).days / 365
    unit_price = round(base_price * 1e6 * (1.035 ** years) / 5e5) * 5e5

    # leasing ~68% (leasing lebih jarang untuk corporate besar)
    is_leasing = rng.random() < (0.55 if cust["corporate"] else 0.75)
    if is_leasing:
        idx = rng.choice(len(LEASING), p=[w for _, w in LEASING])
        leasing = random.choice(LEASING[idx][0])
    else:
        leasing = random.choice(["CASH", "CASH", None, "cash", "TUNAI"])

    # diskon: base + bias salesman + akhir bulan + fleet; cash sedikit lebih tinggi
    disc_pct = base_disc + sm["disc_bias"] + rng.normal(0, 1.0)
    if bdate.day >= 24:
        disc_pct += 1.8
    if cust["corporate"]:
        disc_pct += 0.8
    if not is_leasing:
        disc_pct += 0.6
    if rng.random() < 0.012:  # anomali diskon
        disc_pct = rng.uniform(14, 20)
    disc_pct = float(np.clip(disc_pct, 0.5, 20))
    total_discount = round(unit_price * disc_pct / 100 / 1e5) * 1e5
    price_after = unit_price - total_discount

    # proses: billing -> GI -> physical GI -> BSTK -> STNK
    gi = bdate + timedelta(days=int(rng.integers(0, 4)))
    phys = gi + timedelta(days=int(rng.choice([0, 0, 1, 1, 2, 3, 5, 8])))
    if body not in ("CAB CHASSIS",):  # tunggu karoseri
        phys += timedelta(days=int(rng.integers(0, 10)))
    handover = phys + timedelta(days=int(rng.integers(*ship_days)))
    stnk = bdate + timedelta(days=int(rng.gamma(4, 5.5) + (6 if is_leasing else 0)))
    today = date(2026, 10, 5)
    gi, phys, handover, stnk = [d if d <= today else None for d in (gi, phys, handover, stnk)]
    if bdate > date(2026, 8, 15) and rng.random() < 0.5:
        stnk = None

    term = int(rng.choice([30, 30, 45])) if is_leasing else int(rng.choice([0, 7, 14, 30, 60] if cust["corporate"] else [0, 0, 7, 14]))

    handover_branch = sm["sales_point"] if rng.random() < 0.87 else random.choice(SALES_POINTS)

    plate = None
    if stnk is not None:
        while True:
            plate = f"{plate_code} {rng.integers(8000, 9999)} {random.choice('ABCDEFGHJKLMNOPRSTUVWXYZ')}{random.choice('ABCDEFGHJKLMNOPRSTUVWXYZ')}"
            if plate not in used_plates:
                used_plates.add(plate)
                break

    stnk_name = cust["name"]
    stnk_addr = cust["address"]
    if rng.random() < 0.12:  # STNK atas nama pihak lain
        stnk_name, stnk_addr = fake.name().upper(), fake.street_address().upper()

    rows.append({
        "Tgl Billing": bdate,
        "Spes No": spec,
        "No Customer": cust["id"],
        "Nama Customer": cust["name"],
        "Address": cust["address"],
        "City": cust["city_raw"],
        "Handphone": cust["hp"],
        "Telp Rumah": cust["home"],
        "Telp Kantor": cust["office"],
        "Nama STNK": stnk_name,
        "Alamat STNK": stnk_addr,
        "Chasis": "MHF" + fake.bothify("??#####??#######").upper()[:14],
        "No. Mesin": fam[:2] + fake.bothify("##?-######").upper(),
        "Karoseri Spec.": None if body == "CAB CHASSIS" and rng.random() < 0.5 else body,
        "Color Description": random.choice(COLORS),
        "Actual Physical GI Date": phys,
        "Tgl. GI": gi,
        "Tgl.BSTK": handover,
        "Leasing": leasing,
        "Salesman name": sm["name"],
        "SPV Name": sm["spv"],
        "Sales Point": sm["sales_point"],
        "Package Code": "PKG-FLEET" if cust["corporate"] and rng.random() < 0.5 else random.choice(PACKAGES),
        "Jatuh Tempo": bdate + timedelta(days=term),
        "No Polisi": plate,
        "Harga Unit": unit_price,
        "Harga Discount": price_after,
        "Selesai STNK": stnk,
        "Cabang Serah Terima": handover_branch,
        "Total Discount": total_discount,
    })

df = pd.DataFrame(rows).sort_values("Tgl Billing", kind="stable").reset_index(drop=True)

# ID berurutan mengikuti tanggal billing
seq = df.groupby(df["Tgl Billing"].map(lambda d: (d.year, d.month))).cumcount() + 1
ym = df["Tgl Billing"].map(lambda d: f"{d.year}/{d.month:02d}")
df.insert(0, "No.", range(1, len(df) + 1))
df.insert(1, "Filling Number", [f"FL{2400000 + i:07d}" for i in range(len(df))])
df.insert(2, "No. Faktur", [f"INV/{a}/{b:04d}" for a, b in zip(ym, seq)])
df.insert(5, "No SO", [f"SO{8100000 + i * 3 + int(rng.integers(0, 3)):07d}" for i in range(len(df))])
df.insert(27, "BookNumber", [f"BK-{d.year}-{i:05d}" for i, d in enumerate(df["Tgl Billing"], 1)])
df.insert(30, "No SPK", [f"SPK/{d.year % 100:02d}/{rng.integers(10000, 99999)}" for d in df["Tgl Billing"]])

COLUMNS = ["No.", "Filling Number", "No. Faktur", "Tgl Billing", "Spes No", "No SO", "No Customer",
           "Nama Customer", "Address", "City", "Handphone", "Telp Rumah", "Telp Kantor", "Nama STNK",
           "Alamat STNK", "Chasis", "No. Mesin", "Karoseri Spec.", "Color Description",
           "Actual Physical GI Date", "Tgl. GI", "Tgl.BSTK", "Leasing", "Salesman name", "SPV Name",
           "Sales Point", "Package Code", "BookNumber", "Jatuh Tempo", "No Polisi", "No SPK", "Harga Unit",
           "Harga Discount", "Selesai STNK", "Cabang Serah Terima", "Total Discount"]
df = df[COLUMNS]
assert len(COLUMNS) == 36 and df["No. Faktur"].is_unique

DATE_COLS = ["Tgl Billing", "Actual Physical GI Date", "Tgl. GI", "Tgl.BSTK", "Jatuh Tempo", "Selesai STNK"]
for c in DATE_COLS:
    df[c] = pd.to_datetime(df[c])

OUT.parent.mkdir(parents=True, exist_ok=True)
with pd.ExcelWriter(OUT, engine="openpyxl", datetime_format="DD/MM/YYYY") as xw:
    df.to_excel(xw, sheet_name="Data", index=False)
    ws = xw.sheets["Data"]
    for i, col in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F2A44")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = max(len(col), int(df[col].astype(str).str.len().quantile(0.9))) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
        if col in ("Harga Unit", "Harga Discount", "Total Discount"):
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=i).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

print(f"Saved {len(df)} rows x {df.shape[1]} cols -> {OUT}")
